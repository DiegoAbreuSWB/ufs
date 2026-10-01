"""
make_response.py — Builds the point-by-point response letter (.docx) from
revision/reports/response_content.py, following the structure requested by
IEEE Access: Reviewer's concern / Author response / Author action.
"""

import importlib.util
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))


def load_content():
    spec = importlib.util.spec_from_file_location("response_content", os.path.join(HERE, "reports", "response_content.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def add_par(doc, text="", bold_prefix=None, size=11, italic=False, space_after=6, align=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.size = Pt(size)
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    return p


def add_table(doc, header, rows, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size)
    doc.add_paragraph()


def main():
    c = load_content()
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    add_par(doc, c.MANUSCRIPT_ID, bold_prefix="Original Manuscript ID: ")
    add_par(doc, c.TITLE, bold_prefix="Original Article Title: ")
    add_par(doc, "IEEE Access Editor", bold_prefix="To: ")
    add_par(doc, "Response to reviewers", bold_prefix="Re: ")
    for par in c.COVER:
        add_par(doc, par, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_par(doc, "Best regards,", space_after=0)
    add_par(doc, c.AUTHOR)

    h = doc.add_heading("Summary of the revision", level=1)
    h.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    for par in c.SUMMARY:
        if isinstance(par, dict):
            add_table(doc, par["header"], par["rows"])
        else:
            add_par(doc, par, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    for reviewer in c.REVIEWERS:
        h = doc.add_heading(reviewer["name"], level=1)
        h.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        for par in reviewer.get("intro", []):
            add_par(doc, par, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        for i, item in enumerate(reviewer.get("concerns", []), 1):
            add_par(doc, item["concern"], bold_prefix=f"{reviewer['tag']}, Concern #{i}: ", italic=True,
                    align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            for j, par in enumerate(item["response"]):
                if isinstance(par, dict):
                    add_table(doc, par["header"], par["rows"])
                else:
                    add_par(doc, par, bold_prefix="Author response: " if j == 0 else None, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            for j, par in enumerate(item["action"]):
                add_par(doc, par, bold_prefix="Author action: " if j == 0 else None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                        space_after=12 if j == len(item["action"]) - 1 else 6)

    out = os.path.join(HERE, "reports", "Response_to_Reviewers_ContrastiveXAI_FS_R2.docx")
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()
